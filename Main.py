import PyPDF2
from datetime import datetime
import docx
from pptx import Presentation
from pyexiv2 import Image
from openpyxl import load_workbook
import os


def pdf_metadata_extractor(file_path):
    try:
        with open(file_path, "rb") as pdf_file:
            file_content_object = PyPDF2.PdfReader(pdf_file)
            metadata_info_object = file_content_object.metadata
            print("\n", "-"*50, " Pdf Metadata ", "-"*50)
            for key, value in metadata_info_object.items():
                if key in ["/ModDate", "/CreationDate"]:
                    startindex = value.index(":")+1
                    if '+' in value:
                        endindex = value.index("+")
                    elif 'Z' in value:
                        endindex = value.index("Z")
                    else:
                        endindex = value.index("-")
                    value = value[startindex:endindex]
                    dt_object = datetime.strptime(value, '%Y%m%d%H%M%S')
                    formatted_time = dt_object.strftime('%Y-%m-%d %H:%M:%S')
                    value = formatted_time
                print(f"{key[1:]:30} = {value}")
    except FileNotFoundError:
        print("File is not available at the provided path. Please provide the correct path to the PDF file")
    except BaseException as e:
        print(
            f"Something Went Wrong While Extracting The PDF File Metadata\nError = {e}")


def docx_metadata_extractor(file_path):
    try:
        doc = docx.Document(file_path)
        all_properties = doc.core_properties
        print("\n", "-"*50, " Docx Metadata ", "-"*50)
        for properties in all_properties.__dir__():
            if not properties.startswith("_"):
                print(f"{properties:30} : {getattr(all_properties, properties)}")
    except BaseException as e:
        print(
            f"Something Went Wrong While Extracting The Docx File Metadata\nError = {e}")


def ppt_metadata_extractor(file_path):
    try:
        ppt_file = Presentation(file_path)
        all_properties = ppt_file.core_properties
        print("\n", "-"*50, " PPT Metadata ", "-"*50)
        for properties in all_properties.__dir__():
            if not properties.startswith("_"):
                value = getattr(all_properties, properties)
                # a callable value refers to an object that can be called like a function.
                if callable(value) or properties == "blob":
                    continue
                print(f"{properties:30} : {value}")
    except BaseException as e:
        print(
            f"Something Went Wrong While Extracting The PPT File Metadata\nError = {e}")


def image_metadata_extractor(file_path):
    try:
        image = Image(file_path)
        print("\n", "-"*50, " Image Metadata ", "-"*50)
        exif_data = image.read_exif()
        print("->EXIF metadata:")
        for key, value in exif_data.items():
            if "MakerNote" in str(key):
                continue
            cut = str(key).rindex(".")
            print(f"  {(key[cut+1:]):30}: {value}")
        iptc_data = image.read_iptc()
        print("\n->IPTC metadata:")
        for key, value in iptc_data.items():
            cut = str(key).rindex(".")
            print(f"  {(key[cut+1:]):30}: {value}")
        xmp_data = image.read_xmp()
        print("\n->XMP metadata:")
        for key, value in xmp_data.items():
            cut = str(key).rindex(".")
            print(f"  {(key[cut+1:]):75}: {value}")
    except BaseException as e:
        print(
            f"Something Went Wrong While Extracting The Image File Metadata\nError = {e}")


def text_file_metadata_extractor(file_path):
    try:
        if "\\" in file_path:
            name = file_path[file_path.rfind("\\")+1:]
        else:
            name = file_path
        dt_object_creation_time = datetime.fromtimestamp(
            os.path.getctime(file_path))
        print("\n", "-"*50, " Text File Metadata ", "-"*50)
        dt_object_modification_time = datetime.fromtimestamp(
            os.path.getmtime(file_path))
        print(f"{'MetaData of Text File':30} : {name}")
        print(f"{'File Size':30} : {os.path.getsize(file_path)}")
        print(f"{'Creation Time':30} : {dt_object_creation_time}")
        print(f"{'Last Modification Time':30} : {dt_object_modification_time}")
    except BaseException as e:
        print(
            f"Something Went Wrong While Extracting The Text File Metadata\nError = {e}")


def excel_file_metadata_extractor(file_path):
    try:
        excel_file = load_workbook(file_path)
        all_properties = excel_file.properties
        print("\n", "-"*50, " Excel File Metadata ", "-"*50)
        for prop in all_properties.__dir__():
            if not prop.startswith("_") and prop != "to_tree":
                print(f"{prop:30} : {getattr(all_properties, prop)}")
    except BaseException as e:
        print(
            f"Something Went Wrong While Extracting The Excel File Metadata\nError = {e}")


def main():
    try:
        # Note Please change the path below according to your system.
        # Hard Coded (For testing)
        # files = [r"C:\Users\F.R.I.D.A.Y\Desktop\Metadata Extractor\test\Tkinter-CheatSheet.pdf", r"C:\Users\F.R.I.D.A.Y\Desktop\Metadata Extractor\test\Python Cheat Sheet.docx", r"C:\Users\F.R.I.D.A.Y\Desktop\Metadata Extractor\test\numpy.pptx",
        #          r"C:\Users\F.R.I.D.A.Y\Desktop\Metadata Extractor\test\text.txt", r"C:\Users\F.R.I.D.A.Y\Desktop\Metadata Extractor\test\demo_excel.xlsx", r"C:\Users\F.R.I.D.A.Y\Desktop\Metadata Extractor\test\demo.jpg"]

        # User Input
        files = [input("Enter the path where the file is located: ")]
        metadata_extractors = {
            "pdf": pdf_metadata_extractor,
            "docx": docx_metadata_extractor,
            "pptx": ppt_metadata_extractor,
            "jpg": image_metadata_extractor,
            "jpeg": image_metadata_extractor,
            "txt": text_file_metadata_extractor,
            "xlsx": excel_file_metadata_extractor
        }
        for file_path in files:
            if ":" in file_path:
                file_path = file_path[file_path.index(":")+1:]
            extension = file_path[file_path.rindex(".")+1:].lower()
            if extension in metadata_extractors:
                metadata_extractors[extension](file_path)
            else:
                print("\nThe file path you have provided is invalid or the file is not supported\nThis program only supports pdf, docx, pptx, xlsx, txt, jpg, and jpeg files")
    except KeyboardInterrupt:
        print("Invalid input. Program is terminated.")
        exit(0)
    except BaseException as e:
        print(
            f"Something Went Wrong While Receiving Input From User\nError = {e}")


if __name__ == "__main__":
    main()
